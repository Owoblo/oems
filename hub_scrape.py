"""
Targeted scrape for multi-tenant industrial hub addresses.
Reads hub_analysis.xlsx, searches each base address directly,
and appends only NEW companies to the existing deduped_companies.csv.
"""

import csv
import os
import re
import sys
import time
from datetime import datetime

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PLACES_URL = "https://places.googleapis.com/v1/places:searchText"
FIELD_MASK = "places.id,places.displayName,places.formattedAddress"
REQUEST_DELAY = 0.5
RETRY_DELAYS = [2, 4, 8]

RAW_COLS = [
    "place_id", "company_name", "formatted_address",
    "city_searched", "zone", "keyword_searched", "raw_query",
]
LOG_COLS = [
    "timestamp", "zone", "city", "raw_query",
    "status_code", "number_of_results", "new_results", "error_message",
]

unit_pat = re.compile(
    r'\b(?:unit|suite|ste|bldg|building|bay|lot)\s*[#]?\s*([\w-]+)'
    r'|#\s*([\w-]+)',
    re.IGNORECASE
)


# Matches names that are actually street addresses, not company names
ADDRESS_NAME_PAT = re.compile(
    r'^\d+\s+\w'           # starts with a number then a word (e.g. "550 Trillium")
    r'|^unit\s*[#\d]'      # "Unit #11" / "Unit 4"
    r'|^lot\s*[#\d]'
    r'|^suite\s*[#\d]'
    r'|^bay\s*[#\d]',
    re.IGNORECASE
)

def is_address_name(name: str) -> bool:
    """Return True if the displayName looks like a street address, not a business."""
    return bool(ADDRESS_NAME_PAT.match(str(name).strip()))


def dedup_key(place_id, name, address):
    if place_id:
        return f"id:{place_id}"
    return f"na:{str(name).lower().strip()}|{str(address).lower().strip()}"


def call_api(api_key, query, zone, city):
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": api_key,
        "X-Goog-FieldMask": FIELD_MASK,
    }
    payload = {"textQuery": query, "maxResultCount": 20}

    for attempt, delay in enumerate([0] + RETRY_DELAYS):
        if delay:
            print(f"    Rate-limited. Retry {attempt}/3 in {delay}s…")
            time.sleep(delay)
        try:
            resp = requests.post(PLACES_URL, json=payload, headers=headers, timeout=30)
        except requests.RequestException as exc:
            return 0, [], str(exc)

        if resp.status_code == 429:
            continue
        if resp.status_code != 200:
            return resp.status_code, [], resp.text[:300]

        results = []
        for p in resp.json().get("places", []):
            results.append({
                "place_id": p.get("id", ""),
                "company_name": (p.get("displayName") or {}).get("text", ""),
                "formatted_address": p.get("formattedAddress", ""),
                "city_searched": city,
                "zone": zone,
                "keyword_searched": "hub_address_search",
                "raw_query": query,
            })
        return resp.status_code, results, ""

    return 429, [], "Rate-limited after all retries"


def run(output_dir):
    api_key = os.environ.get("GOOGLE_PLACES_API_KEY", "").strip()
    if not api_key:
        print("ERROR: GOOGLE_PLACES_API_KEY not set.")
        sys.exit(1)

    hub_xlsx   = os.path.join(output_dir, "hub_analysis.xlsx")
    raw_path   = os.path.join(output_dir, "raw_results.csv")
    dedup_path = os.path.join(output_dir, "deduped_companies.csv")
    log_path   = os.path.join(output_dir, "hub_run_log.csv")
    new_path   = os.path.join(output_dir, "hub_new_companies.csv")

    # Load all hubs
    hubs = pd.read_excel(hub_xlsx, sheet_name="All Potential Hubs")
    print(f"Loaded {len(hubs)} hub addresses from hub_analysis.xlsx")

    # Load existing deduped keys so we don't re-add anything
    existing = pd.read_csv(dedup_path)
    seen_keys = set()
    for _, row in existing.iterrows():
        seen_keys.add(dedup_key(row["place_id"], row["company_name"], row["formatted_address"]))
    print(f"Loaded {len(seen_keys)} existing companies to skip duplicates")
    print(f"Running {len(hubs)} targeted address searches…\n")

    # Write log header
    with open(log_path, "w", newline="", encoding="utf-8") as f:
        csv.DictWriter(f, fieldnames=LOG_COLS).writeheader()

    # Write new companies header
    with open(new_path, "w", newline="", encoding="utf-8") as f:
        csv.DictWriter(f, fieldnames=RAW_COLS).writeheader()

    total_new = 0
    total_raw = 0

    for i, hub_row in enumerate(hubs.itertuples(), 1):
        base_addr = hub_row.base_address
        city      = hub_row.city
        zone      = hub_row.zone

        # Build a clean query: just the address — Google surfaces tenants this way
        query = base_addr.strip().rstrip(",").strip()

        print(f"[{i}/{len(hubs)}] {query}")

        ts = datetime.utcnow().isoformat(timespec="seconds") + "Z"
        status_code, results, error_msg = call_api(api_key, query, zone, city)

        new_results = []
        if results:
            total_raw += len(results)
            # Append all to raw
            with open(raw_path, "a", newline="", encoding="utf-8") as f:
                csv.DictWriter(f, fieldnames=RAW_COLS, extrasaction="ignore").writerows(results)

            # Filter to genuinely new companies (skip address-as-name junk records)
            for r in results:
                if is_address_name(r["company_name"]):
                    continue
                key = dedup_key(r["place_id"], r["company_name"], r["formatted_address"])
                if key not in seen_keys:
                    seen_keys.add(key)
                    new_results.append(r)

            if new_results:
                total_new += len(new_results)
                with open(dedup_path, "a", newline="", encoding="utf-8") as f:
                    csv.DictWriter(f, fieldnames=RAW_COLS, extrasaction="ignore").writerows(new_results)
                with open(new_path, "a", newline="", encoding="utf-8") as f:
                    csv.DictWriter(f, fieldnames=RAW_COLS, extrasaction="ignore").writerows(new_results)

        status_str = f"→ {len(results)} results, {len(new_results)} new" if not error_msg else f"ERROR: {error_msg}"
        print(f"    {status_str}")

        with open(log_path, "a", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=LOG_COLS, extrasaction="ignore").writerow({
                "timestamp": ts, "zone": zone, "city": city,
                "raw_query": query, "status_code": status_code,
                "number_of_results": len(results),
                "new_results": len(new_results),
                "error_message": error_msg,
            })

        time.sleep(REQUEST_DELAY)

    print(f"\n{'='*60}")
    print(f"  Hub searches run  : {len(hubs)}")
    print(f"  Raw results       : {total_raw}")
    print(f"  NEW companies     : {total_new}")
    print(f"  Output dir        : {output_dir}")
    print(f"{'='*60}")

    if total_new > 0:
        print("\nExporting updated XLSX…")
        export_xlsx(output_dir)
        print("  Done.")


def export_xlsx(output_dir):
    csv_path  = os.path.join(output_dir, "deduped_companies.csv")
    xlsx_path = os.path.join(output_dir, "deduped_companies.xlsx")
    df = pd.read_csv(csv_path)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    col_widths = {
        "place_id": 30, "company_name": 45, "formatted_address": 55,
        "city_searched": 20, "zone": 42, "keyword_searched": 28, "raw_query": 55,
    }

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Leads", index=False)
        ws = writer.sheets["All Leads"]
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = col_widths.get(col, 20)
        ws.freeze_panes = "A2"

        summary = (
            df.groupby("zone")
            .agg(unique_companies=("place_id", "nunique"), cities=("city_searched", "nunique"))
            .reset_index()
        )
        summary.columns = ["Zone", "Unique Companies", "Cities Covered"]
        summary.to_excel(writer, sheet_name="Summary", index=False)
        ws2 = writer.sheets["Summary"]
        for cell in ws2[1]:
            cell.fill = header_fill; cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        ws2.column_dimensions["A"].width = 45
        ws2.column_dimensions["B"].width = 22
        ws2.column_dimensions["C"].width = 18
        ws2.freeze_panes = "A2"

        for zone_name, zone_df in df.groupby("zone"):
            short = zone_name.split("—")[-1].strip().replace("/", "-")[:28]
            zone_df.to_excel(writer, sheet_name=short, index=False)
            ws3 = writer.sheets[short]
            for cell in ws3[1]:
                cell.fill = header_fill; cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for i, col in enumerate(zone_df.columns, 1):
                ws3.column_dimensions[get_column_letter(i)].width = col_widths.get(col, 20)
            ws3.freeze_panes = "A2"


if __name__ == "__main__":
    run("./places_output")
